"""空白招标 Word 模板的解析与 OCR 指标锚点定位。

模板不是必需输入。提供后只提取结构化的章节、表格标签和显式指标名，
不把模板全文或原始文本复制进报告；OCR 阶段使用这些标签做保守的空间配对，
无法配对时仍回退到通用 OCR 候选并保留人工复核状态。
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from .canon import write_json
from .normalize import to_halfwidth


class TemplateDependencyError(RuntimeError):
    """模板已被选中，但当前运行环境没有解析 DOCX 所需依赖。"""


_GENERIC_LABELS = {
    "指标", "指标名称", "参数", "技术参数", "技术指标", "要求", "技术要求",
    "响应", "响应情况", "投标响应", "单位", "备注", "序号", "项目", "项目名称",
    "采购需求", "技术需求", "符合性", "是否响应",
}
_DISCOVERY_TOKENS = ("招标", "采购", "模板", "技术规范", "技术参数", "评分", "需求")
_NUMBER_PREFIX = re.compile(r"^\s*(?:[（(]?[一二三四五六七八九十百\d]+[）).、]|[-—•●])\s*")
_SPACE_RE = re.compile(r"\s+")


@dataclass
class TemplateMetric:
    metric_id: str
    label: str
    normalized_label: str
    section: str | None = None
    unit: str | None = None
    source_kind: str = "docx_table"
    source_location: dict[str, Any] = field(default_factory=dict)


@dataclass
class TemplateResolution:
    status: str
    path: Path | None = None
    relative_path: str | None = None
    candidates: list[str] = field(default_factory=list)
    message: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "path": str(self.path) if self.path else None,
            "relative_path": self.relative_path,
            "candidates": self.candidates,
            "message": self.message,
        }


@dataclass
class TemplateSpec:
    status: str
    mode: str
    source_path: str | None = None
    source_sha256: str | None = None
    metrics: list[TemplateMetric] = field(default_factory=list)
    sections: list[str] = field(default_factory=list)
    candidates: list[str] = field(default_factory=list)
    message: str = ""

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["metrics"] = [asdict(metric) for metric in self.metrics]
        return data


def normalize_template_label(text: str) -> str:
    """把模板标签规整为可跨 OCR 版式比较的短键。"""
    value = to_halfwidth(str(text or "")).strip().lower()
    value = _NUMBER_PREFIX.sub("", value)
    value = re.sub(r"[：:；;，,。.!！?？（）()【】\[\]{}<>《》/\\|_—–-]", "", value)
    return _SPACE_RE.sub("", value)


def _relative_path(root: Path, path: Path) -> str:
    return path.resolve().relative_to(root.resolve()).as_posix()


def _is_inside(root: Path, path: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def _candidate_score(path: Path) -> int:
    name = path.stem
    return sum(3 if token in name else 0 for token in _DISCOVERY_TOKENS) + (1 if "procurement" in path.parts else 0)


def _looks_like_template_filename(path: Path) -> bool:
    name = path.stem
    return any(token in name for token in ("模板", "空白", "template"))


def resolve_template(project_dir: Path, configured_path: str | None = None) -> TemplateResolution:
    """解析 project.yaml 指定的模板；未指定时仅在 procurement/ 做保守发现。

    自动发现只在只有一个高相关候选时生效，多个候选不猜，避免把正式招标文件
    或历史版本误当成空白模板。
    """
    root = project_dir.resolve()
    if configured_path:
        raw = Path(configured_path).expanduser()
        path = raw if raw.is_absolute() else root / raw
        path = path.resolve()
        if not _is_inside(root, path):
            return TemplateResolution(
                status="invalid",
                message="tender_template_path 必须位于项目目录内",
            )
        if path.suffix.lower() != ".docx":
            return TemplateResolution(status="invalid", message="招标模板必须是 .docx Word 文件")
        if not path.exists() or not path.is_file():
            return TemplateResolution(
                status="missing",
                relative_path=_relative_path(root, path) if _is_inside(root, path) else None,
                message=f"未找到配置的 Word 模板：{path}",
            )
        return TemplateResolution(
            status="provided",
            path=path,
            relative_path=_relative_path(root, path),
            message="已指定空白招标 Word 模板；模板优先定位可大幅提高指标识别准确度",
        )

    procurement = root / "procurement"
    if not procurement.exists():
        return TemplateResolution(
            status="not_provided",
            message="未提供空白招标 Word 模板；继续使用通用 OCR 兜底。提供模板可大幅提高指标定位和识别准确度",
        )
    candidates = sorted(
        (p for p in procurement.rglob("*.docx") if p.is_file() and not p.name.startswith("~$")),
        key=lambda p: (-_candidate_score(p), p.as_posix()),
    )
    scored = [p for p in candidates if _candidate_score(p) >= 4 and _looks_like_template_filename(p)]
    if len(scored) == 1:
        path = scored[0].resolve()
        return TemplateResolution(
            status="provided",
            path=path,
            relative_path=_relative_path(root, path),
            candidates=[_relative_path(root, p) for p in scored],
            message="已从 procurement/ 唯一识别 Word 模板；模板优先定位可大幅提高指标识别准确度",
        )
    if len(scored) > 1:
        return TemplateResolution(
            status="ambiguous",
            candidates=[_relative_path(root, p) for p in scored],
            message="发现多个可能的 Word 模板，请在 project.yaml 的 tender_template_path 中明确指定；未指定时使用通用 OCR 兜底",
        )
    return TemplateResolution(
        status="not_provided",
        message="未提供空白招标 Word 模板；继续使用通用 OCR 兜底。提供模板可大幅提高指标定位和识别准确度",
    )


def _clean_cell(value: str) -> str:
    return _SPACE_RE.sub(" ", str(value or "").replace("\n", " ")).strip()


def _looks_like_metric_label(value: str) -> bool:
    normalized = normalize_template_label(value)
    if not normalized or normalized in {normalize_template_label(x) for x in _GENERIC_LABELS}:
        return False
    if len(normalized) < 2 or len(normalized) > 80:
        return False
    if re.fullmatch(r"[\d.]+", normalized):
        return False
    return any("指标" in normalized or token in normalized for token in ("参数", "规格", "型号", "性能", "容量", "尺寸", "材质", "功率", "数量", "等级", "标准", "要求"))


def _extract_unit(label: str, cells: list[str]) -> tuple[str, str | None]:
    match = re.search(r"[（(]([^（）()]{1,12})[）)]\s*$", label)
    if match:
        return label[:match.start()].strip(), match.group(1).strip()
    for cell in cells[1:]:
        if cell and len(cell) <= 12 and any(ch in cell for ch in ("米", "mm", "cm", "kg", "kw", "%", "个", "套", "台", "级")):
            return label, cell
    return label, None


def _metric_id(label: str, section: str | None, location: dict[str, Any]) -> str:
    raw = "|".join((normalize_template_label(label), section or "", repr(sorted(location.items()))))
    return "M-" + hashlib.sha256(raw.encode("utf-8")).hexdigest()[:10]


def parse_template(path: Path) -> TemplateSpec:
    """从 DOCX 提取章节和指标锚点；不保存模板原文。"""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - 环境相关分支
        raise TemplateDependencyError("解析 Word 模板需要 python-docx；请按预检提示安装") from exc

    path = path.resolve()
    document = Document(path)
    sections: list[str] = []
    metrics: list[TemplateMetric] = []
    seen: set[str] = set()
    current_section: str | None = None

    for index, paragraph in enumerate(document.paragraphs):
        text = _clean_cell(paragraph.text)
        if not text:
            continue
        style = str(getattr(paragraph.style, "name", "") or "").lower()
        if style.startswith("heading") or style in {"title", "标题", "副标题"}:
            current_section = text
            if text not in sections:
                sections.append(text)
        match = re.match(r"^(.{2,80}?)[：:]\s*(?:____+|（待填|待填写|填写|响应|符合|是|否)?\s*$", text)
        if match:
            label, unit = _extract_unit(match.group(1).strip(), [])
            if _looks_like_metric_label(label):
                location = {"kind": "docx_body", "index": index}
                normalized = normalize_template_label(label)
                if normalized not in seen:
                    seen.add(normalized)
                    metrics.append(TemplateMetric(_metric_id(label, current_section, location), label, normalized, current_section, unit, "docx_body", location))

    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            cells = [_clean_cell(cell.text) for cell in row.cells]
            nonempty = [cell for cell in cells if cell]
            if not nonempty:
                continue
            label, unit = _extract_unit(nonempty[0], nonempty)
            if not _looks_like_metric_label(label):
                continue
            location = {"kind": "docx_table", "table": table_index, "row": row_index, "col": 0}
            normalized = normalize_template_label(label)
            if normalized in seen:
                continue
            seen.add(normalized)
            metrics.append(TemplateMetric(_metric_id(label, current_section, location), label, normalized, current_section, unit, "docx_table", location))

    source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    return TemplateSpec(
        status="parsed",
        mode="template_aligned" if metrics else "template_empty",
        source_path=path.as_posix(),
        source_sha256=source_sha256,
        metrics=metrics,
        sections=sections,
        message=("已解析模板指标锚点；将优先按模板标签定位，无法配对时回退通用 OCR"
                  if metrics else "已读取模板但未发现可用指标锚点；将继续使用通用 OCR 兜底"),
    )


def load_template_spec(project_dir: Path, configured_path: str | None = None) -> TemplateSpec:
    resolution = resolve_template(project_dir, configured_path)
    if resolution.status != "provided" or not resolution.path:
        return TemplateSpec(
            status=resolution.status,
            mode="generic_ocr",
            source_path=resolution.relative_path,
            candidates=resolution.candidates,
            message=resolution.message,
        )
    try:
        spec = parse_template(resolution.path)
    except TemplateDependencyError as exc:
        return TemplateSpec(
            status="dependency_missing",
            mode="generic_ocr",
            source_path=resolution.relative_path,
            candidates=resolution.candidates,
            message=str(exc),
        )
    spec.source_path = resolution.relative_path
    spec.candidates = resolution.candidates
    return spec


def write_template_spec(project_dir: Path, spec: TemplateSpec) -> Path:
    path = project_dir / "output" / "interim" / "template-spec.json"
    write_json(path, spec.to_dict())
    return path
