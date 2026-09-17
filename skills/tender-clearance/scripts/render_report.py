#!/usr/bin/env python3
"""阶段六：渲染清标报告与结构化结果。

生成（方案 7）：
- output/清标结果.json     全量结构化结果（验收基准）
- output/清标报告.md       审阅版 Markdown（验收基准）
- output/证据索引.csv      证据登记表（utf-8-sig）
- output/人工复核清单.csv  I 级/主体歧义/低置信度/缺口（utf-8-sig）
- output/清标报告.docx     展示副本（python-docx 生成；可选）

渲染前运行完整性校验；`redaction_mode: standard` 额外执行完整身份证号/手机号扫描，
`none` 模式按项目授权保留原值，但任一模式都禁止输出访问凭据。
"""

from __future__ import annotations

import csv
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import typer

from tc.canon import load_json, write_json
from tc.models import (
    AlertLevel,
    ContentFile,
    EntitiesFile,
    Evidence,
    EvidenceFile,
    ExternalEvidenceFile,
    FindingsFile,
    Finding,
    InventoryFile,
    MatchesFile,
    MetadataFile,
    ReportBundle,
    ROLE_LABELS,
    SOURCE_LABELS,
)
from tc.projio import ProjectError, ensure_output_dirs, load_project_config
from tc.srm_gate import SrmReportGateError, validate_srm_report_gate
from tc.normalize import normalize_company_name
from tc.process_artifacts import build_process_artifacts, ocr_completion, validate_report_input
from tc.field_decisions import decisions_for_fields

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "清标报告.md.jinja"

# 边界用字母数字整体排除：避免把哈希/摘要中的数字串误判为手机号或证件号
_ID18_ANY = re.compile(r"(?<![0-9A-Za-z])\d{6}(?:19|20)\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\d{3}[\dXx](?![0-9A-Za-z])")
_MOBILE_ANY = re.compile(r"(?<![0-9A-Za-z])1[3-9]\d{9}(?![0-9A-Za-z])")

app = typer.Typer(help="渲染清标报告与结构化结果")


@app.command()
def run(
    project_dir: Path = typer.Argument(..., exists=True, file_okay=False, help="项目目录"),
    profile: str = typer.Option("report", help="输出档位：report / review / workpaper"),
    docx: bool = typer.Option(False, "--docx/--no-docx", help="兼容选项：额外生成 DOCX"),
    skip_srm_gate: bool = typer.Option(False, "--skip-srm-gate", help="仅供离线测试夹具使用，正式报告不得跳过"),
) -> None:
    if profile not in {"report", "review", "workpaper"}:
        typer.secho(f"[错误] 不支持的输出档位：{profile}", fg=typer.colors.RED, err=True)
        raise typer.Exit(code=2)
    want_docx = docx or profile in {"review", "workpaper"}
    try:
        cfg = load_project_config(project_dir)
        interim = project_dir / "output/interim"
        inventory = InventoryFile(**load_json(interim / "inventory.json"))
        entities = EntitiesFile(**load_json(interim / "entities.json"))
        matches = MatchesFile(**load_json(interim / "matches.json"))
        meta = MetadataFile(**load_json(interim / "metadata.json"))
        ext = ExternalEvidenceFile(**load_json(interim / "external.json"))
        findings = FindingsFile(**load_json(interim / "findings.json"))
        content = ContentFile(**load_json(interim / "content.json"))
        bid_analysis = load_json(interim / "bid-analysis.json") if (interim / "bid-analysis.json").exists() else {}
    except (ProjectError, FileNotFoundError) as exc:
        typer.secho(f"[错误] {exc}", fg=typer.colors.RED, err=True)
        raise typer.Exit(code=2)
    if not skip_srm_gate:
        try:
            validate_srm_report_gate(entities, ext)
        except SrmReportGateError as exc:
            typer.secho(f"[错误] {exc}", fg=typer.colors.RED, err=True)
            raise typer.Exit(code=4)

    # 生成过程底稿和报告安全输入。二者都写入本地项目目录；渲染器继续直接
    # 使用结构化中间产物，不把 OCR 原文、块坐标或数值置信度交给宿主对话。
    build_process_artifacts(
        project_dir,
        inventory=inventory,
        content=content,
        entities=entities,
        findings=findings,
    )
    report_input = load_json(interim / "report-input.json")
    forbidden = validate_report_input(report_input)
    if forbidden:
        typer.secho(
            f"[错误] report-input.json 含禁止进入报告上下文的字段：{', '.join(forbidden)}",
            fg=typer.colors.RED,
            err=True,
        )
        raise typer.Exit(code=6)
    completion = ocr_completion(project_dir)
    if not skip_srm_gate and not completion["terminal"]:
        typer.secho(
            f"[错误] OCR 尚未全部完成：仍有 {len(completion['pending_job_ids'])} 个任务没有终态；"
            "正式报告已阻断，请先补齐 OCR 结果或提交明确的失败/阻断结果",
            fg=typer.colors.RED,
            err=True,
        )
        raise typer.Exit(code=5)
    out_dir, interim = ensure_output_dirs(project_dir)

    evidence: dict[str, Evidence] = {}
    for name in ("evidence-content.json", "evidence-metadata.json", "evidence-external.json",
                 "evidence-external-queries.json"):
        p = interim / name
        if p.exists():
            for e in load_json(p).get("evidence", []):
                ev = Evidence(**e)
                evidence[ev.evidence_id] = ev
    low_conf_ids = set()
    lc = interim / "low_confidence.json"
    if lc.exists():
        low_conf_ids = set(load_json(lc).get("evidence_ids", []))

    bundle = ReportBundle(
        run=inventory.run, project=cfg, inventory=inventory, entities=entities,
        matches=matches, metadata=meta, external=ext, findings=findings,
        bid_analysis=bid_analysis,
        evidence=EvidenceFile(run=inventory.run, evidence=sorted(evidence.values(), key=lambda x: x.evidence_id)),
    )
    result_path = out_dir / "清标结果.json"
    write_json(result_path, bundle.model_dump(mode="json"))

    ctx = _build_context(cfg, inventory, entities, meta, ext, findings, evidence, low_conf_ids, content, project_dir, bid_analysis, report_input)
    md_path = out_dir / "清标报告.md"
    _render_markdown(ctx, md_path)

    # PDF 审阅版（与 Markdown 同源同内容；失败不影响验收基准输出）
    pdf_path = out_dir / "清标报告.pdf"
    try:
        from tc.md2pdf import md_to_pdf

        md_to_pdf(md_path.read_text(encoding="utf-8"), pdf_path, base_dir=out_dir)
    except Exception as exc:  # noqa: BLE001
        typer.secho(f"[提示] PDF 生成失败（不影响验收基准输出）：{exc}", fg=typer.colors.YELLOW)

    _render_evidence_csv(evidence, inventory, out_dir / "证据索引.csv")
    _render_review_csv(
        findings, entities, low_conf_ids, inventory, content,
        out_dir / "人工复核清单.csv",
        field_decisions=report_input.get("field_decisions", []),
    )
    if want_docx:
        try:
            _render_docx(ctx, out_dir / "清标报告.docx")
        except Exception as exc:  # noqa: BLE001
            typer.secho(f"[提示] DOCX 生成失败（不影响验收基准输出）：{exc}", fg=typer.colors.YELLOW)

    if profile == "workpaper":
        worksheet = Path(__file__).with_name("render_worksheet.py")
        proc = subprocess.run([sys.executable, str(worksheet), str(project_dir)], capture_output=True, text=True)
        if proc.returncode:
            typer.secho(f"[提示] Excel 工作底稿生成失败（不影响报告）：{(proc.stdout + proc.stderr)[-600:]}", fg=typer.colors.YELLOW)

    # standard 模式下执行脱敏检查；none 模式按用户授权输出完整身份证号/手机号。
    check_paths = [result_path, md_path, out_dir / "证据索引.csv", out_dir / "人工复核清单.csv"]
    if pdf_path.exists():
        check_paths.append(pdf_path)
    if cfg.redaction_mode == "standard":
        for path in check_paths:
            if path.suffix == ".pdf":
                import fitz

                text = "\n".join(page.get_text() for page in fitz.open(path))
            else:
                text = path.read_text(encoding="utf-8")
            hits = _find_sensitive(text)
            if hits:
                typer.secho(f"[错误] 脱敏检查失败：{path.name} 中发现疑似完整证件号/手机号：{hits}", fg=typer.colors.RED, err=True)
                raise typer.Exit(code=3)

    exports = out_dir / "exports" / profile
    exports.mkdir(parents=True, exist_ok=True)
    export_names = ["清标结果.json", "清标报告.md", "证据索引.csv", "人工复核清单.csv"]
    if pdf_path.exists():
        export_names.append("清标报告.pdf")
    if want_docx and (out_dir / "清标报告.docx").exists():
        export_names.append("清标报告.docx")
    if profile == "workpaper" and (out_dir / "清标底稿.xlsx").exists():
        export_names.append("清标底稿.xlsx")
    for name in export_names:
        shutil.copy2(out_dir / name, exports / name)

    typer.secho(
        f"[OK] 报告已生成 → {out_dir / '清标报告.md'}、清标结果.json、证据索引.csv、人工复核清单.csv"
        + ("、清标报告.pdf" if pdf_path.exists() else "")
        + f"；档位={profile} → {exports}",
        fg=typer.colors.GREEN,
    )


def _find_sensitive(text: str) -> list[str]:
    hits = [m.group(0) for m in _MOBILE_ANY.finditer(text)]
    for m in _ID18_ANY.finditer(text):
        hits.append(m.group(0))
    return hits[:5]


# --------------------------------------------------------------------- 上下文


def _sup_name(entities: EntitiesFile) -> dict[str, str]:
    return {s.supplier_id: _display_company_name(s.display_name) for s in entities.suppliers}


_LEVEL_LABEL: dict[AlertLevel, str] = {"I": "I 级（高）", "II": "II 级（中）", "III": "III 级（低）"}
_STATUS_LABEL = {
    "open": "待复核",
    "human_review_required": "人工复核必需",
    "resolved": "已复核",
}
_DOMAIN_LABEL = {
    "identity": "主体身份",
    "metadata": "文件属性",
    "ownership": "股权关联",
    "judicial": "司法经营",
    "dishonesty": "违法失信",
    "coverage": "查询覆盖",
}

_QUERY_STATUS_LABEL = {
    "match": "已查询到匹配记录",
    "no_match_verified": "已查询，未发现匹配记录",
    "no_result": "已查询，但结果无法判定",
    "not_queried": "未查询",
    "blocked": "访问受阻",
    "failed": "查询失败",
    "needs_manual_review": "待人工复核",
}


def _display_company_name(value: str | None) -> str:
    """报告展示名：去除扫描件中的盖章注记，原始证据不做改写。"""
    if not value:
        return ""
    return re.sub(r"\s*[（(](?:公章|盖章|签章)[）)]\s*$", "", str(value)).strip()


def _query_status_label(value: str | None) -> str:
    return _QUERY_STATUS_LABEL.get(str(value or ""), "未取得")


def _frow(f: Finding, names: dict[str, str]) -> dict:
    return {
        "finding_id": f.finding_id,
        "rule_id": f.rule_id,
        "domain": f.domain,
        "domain_label": _DOMAIN_LABEL.get(f.domain, f.domain),
        "level": f.level,
        "level_label": _LEVEL_LABEL[f.level],
        "evidence_strength": f.evidence_strength,
        "supplier_ids": f.supplier_ids,
        "suppliers_label": "、".join(names.get(s, s) for s in f.supplier_ids) or "（未归属供应商）",
        "fact": f.fact,
        "status": f.status,
        "status_label": _STATUS_LABEL[f.status],
        "recommendation": f.recommendation,
        "clause_note": f.procurement_clause_note,
        "evidence_ids": f.evidence_ids,
        "related_ids": f.related_ids,
        "subject_confirmation": f.subject_confirmation,
    }


def _first_field_value(fields: dict[str, object], keys: tuple[str, ...]) -> str:
    """读取外部画像中的第一个有效字段值，不把占位符当作事实。"""
    placeholders = {"", "-", "--", "暂无", "无", "未返回", "未取得"}
    for key in keys:
        value = str(fields.get(key) or "").strip()
        if value not in placeholders:
            return value
    return ""


def _cover_tenderer_candidates(content: ContentFile, inventory: InventoryFile,
                               suppliers: list) -> list[str]:
    """从商务标首页的可追溯企业名候选中识别招标人。

    首页通常同时出现招标人和投标人；只接受商务标首页、非低置信度、
    可完整匹配公司名称的候选，并排除已归组的投标人。没有足够证据时
    返回空列表，由报告保留“未取得”，不把投标人名称冒充招标人。
    """
    business_docs = {
        d.document_id for d in inventory.documents
        if d.category == "bid" and d.bid_subtype == "business"
    }
    def subject_key(value: str) -> str:
        value = re.sub(r"[（(](?:公章|盖章|签章)[）)]$", "", value.strip())
        return normalize_company_name(value)

    supplier_keys = {subject_key(s.declared_name or s.display_name) for s in suppliers}
    exact_company = re.compile(
        r"^[一-龥A-Za-z0-9（）()]{4,60}(?:股份有限公司|有限责任公司|有限公司|集团公司)$"
    )
    counts: dict[str, int] = {}
    for field in content.fields:
        if (
            field.field != "company_name"
            or field.document_id not in business_docs
            or field.location.get("page") != 1
            or not field.location.get("cover")
            or field.low_confidence
        ):
            continue
        value = str(field.normalized or field.value_masked or "").strip()
        if not exact_company.fullmatch(value):
            continue
        if subject_key(value) in supplier_keys:
            continue
        counts[value] = counts.get(value, 0) + 1
    # 多个招标人候选不得全部拼接进正式字段；交由字段决策层人工复核。
    return [value for value, _count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))] if len(counts) == 1 else []


def _build_context(cfg, inventory, entities, meta, ext, findings_file, evidence, low_conf_ids, content, project_dir: Path, bid_analysis: dict | None = None, report_input: dict | None = None) -> dict:
    names = _sup_name(entities)
    sups = entities.suppliers
    findings = findings_file.findings
    quote_by_supplier = {str(x.get("supplier_id")): x for x in (bid_analysis or {}).get("suppliers", [])}
    # 新口径：公共项只接受封面第一页；主体信息只接受商务标。
    project_values: dict[str, list[str]] = {k: [] for k in ("tenderer", "project_name", "project_code", "bid_date")}
    decision_rows = (report_input or {}).get("field_decisions", [])
    if not decision_rows:
        document_groups: dict[str, str] = {}
        grouping_path = project_dir / "output/interim/supplier-grouping.json"
        if grouping_path.exists():
            try:
                grouping = load_json(grouping_path)
                document_groups = {
                    str(item["document_id"]): str(item["group_id"])
                    for item in grouping.get("documents", [])
                    if item.get("status") == "assigned" and item.get("group_id")
                }
            except (OSError, ValueError, TypeError):
                document_groups = {}
        decision_rows = [d.as_dict() for d in decisions_for_fields(
            content.fields, document_groups=document_groups
        )]
    for decision in decision_rows:
        if (
            decision.get("scope") == "project"
            and decision.get("field") in project_values
            and decision.get("status") == "selected"
            and decision.get("value")
        ):
            project_values[decision["field"]].append(str(decision["value"]))
    project_identity = {k: sorted(set(v)) for k, v in project_values.items()}
    # 商务标首页常以标题直接出现招标人，没有“招标人：”标签；在已提取
    # 公共字段为空时，使用同页公司名候选补齐，但不能把投标人冒充招标人。
    if not project_identity["tenderer"]:
        project_identity["tenderer"] = _cover_tenderer_candidates(content, inventory, sups)
    # project.yaml 的 bid_deadline 是用户在流程开始前确认的截止时间。报告
    # 不把截止时间伪装成某家公司的“投标日期”，统一用明确的合并口径展示。
    if not project_identity["bid_date"] and cfg.bid_deadline is not None:
        deadline_date = cfg.bid_deadline.date().isoformat() if hasattr(cfg.bid_deadline, "date") else str(cfg.bid_deadline)
        project_identity["bid_date"] = [f"{deadline_date}（投标截止时间）"]
    # SRM 企业画像是公司代码、法人信息的优先来源；商务标只负责授权代表。
    # 这里按供应商聚合 registration，避免把 SRM 查询结果漏在第 3 节而不回填第 2 节。
    srm_registration_by_supplier: dict[str, dict[str, object]] = {}
    for record in ext.records:
        if record.record_kind != "registration" or not record.supplier_id or not record.fields:
            continue
        previous = srm_registration_by_supplier.get(record.supplier_id)
        if previous is None or len(record.fields) > len(previous["fields"]):
            srm_registration_by_supplier[record.supplier_id] = {
                "fields": dict(record.fields),
                "evidence_ids": list(record.evidence_ids),
                "subject_confirmation": record.subject_confirmation,
            }

    business_rows = []
    for s in sups:
        def party(role: str, attr: str) -> str:
            vals = [getattr(p, attr) for p in entities.parties if p.supplier_id == s.supplier_id and p.role == role]
            values = sorted({str(v) for v in vals if v})
            if len(values) == 1:
                return values[0]
            if len(values) > 1:
                return "待人工复核（多个候选）"
            return "未取得"
        srm_identity = srm_registration_by_supplier.get(s.supplier_id, {})
        srm_fields = srm_identity.get("fields", {}) if isinstance(srm_identity, dict) else {}
        if not isinstance(srm_fields, dict):
            srm_fields = {}
        srm_uscc = _first_field_value(srm_fields, ("统一社会信用代码", "信用代码"))
        srm_legal_name = _first_field_value(srm_fields, ("法定代表人", "法人代表", "法人"))
        srm_legal_id = _first_field_value(
            srm_fields,
            ("法定代表人身份证号", "法人代表身份证号", "法人身份证号", "法定代表人证件号"),
        )
        business_rows.append({
            "supplier": names.get(s.supplier_id, _display_company_name(s.display_name)),
            "company": _display_company_name(s.declared_name) or "未取得",
            "uscc": srm_uscc or s.uscc or ("待人工复核（多个候选）" if len(s.uscc_candidates) > 1 else "未取得"),
            "legal_name": srm_legal_name or party("legal_rep", "name"),
            "legal_id": srm_legal_id or party("legal_rep", "id_mask"),
            "agent_name": party("bid_agent", "name"),
            "agent_id": party("bid_agent", "id_mask"),
            "identity_source": "富奥 SRM 企业画像" if srm_identity else "标书商务标",
            "agent_source": "标书商务标扫描页",
        })

    srm_queries = {
        s.supplier_id: next(
            (q for q in ext.queries if q.source_id == "srm" and q.subject_supplier_id == s.supplier_id),
            None,
        )
        for s in sups
    }

    def _srm_query_status(supplier_id: str) -> str:
        q = srm_queries.get(supplier_id)
        return q.status if q else "未查询"

    srm_registration = []
    for supplier in sups:
        rows = [r for r in ext.records if r.record_kind == "registration" and r.supplier_id == supplier.supplier_id]
        if rows:
            for r in rows:
                srm_registration.append({
                    "supplier": names.get(supplier.supplier_id, _display_company_name(supplier.display_name)),
                    "fields": r.fields,
                    "evidence_ids": r.evidence_ids,
                    "query_status": _srm_query_status(supplier.supplier_id),
                })
        else:
            srm_registration.append({
                "supplier": names.get(supplier.supplier_id, _display_company_name(supplier.display_name)),
                "fields": {},
                "evidence_ids": [],
                "query_status": _srm_query_status(supplier.supplier_id),
            })

    ownership_rows_raw = []
    for o in ext.ownership:
        ownership_rows_raw.append({
            "supplier": names.get(o.supplier_id or "", "未归属"),
            "shareholder": o.from_party_name or o.from_company_name or "未取得",
            "amount": "未返回" if o.share_ratio is None else str(o.share_ratio),
            "time": str(o.relation_date) if o.relation_date else "未返回",
            "ratio": str(o.share_ratio) if o.share_ratio is not None else "未返回",
            "method": "未返回", "evidence_ids": o.evidence_ids,
        })
    for r in ext.records:
        if r.record_kind == "ownership":
            fields = r.fields
            ownership_rows_raw.append({
                "supplier": names.get(r.supplier_id or "", r.subject_name or "未归属"),
                "shareholder": _first_field_value(fields, (
                    "股东名称", "股东", "shareholder", "公司名称或股东名称", "发起人名称",
                )) or "未取得",
                "amount": _first_field_value(fields, ("投资金额", "出资额", "认缴出资额")) or "未返回",
                "time": _first_field_value(fields, ("认缴时间", "认缴日期")) or "未返回",
                "ratio": _first_field_value(fields, ("认缴比例", "持股比例")) or "未返回",
                "method": _first_field_value(fields, ("认缴出资方式", "出资方式")) or "未返回",
                "evidence_ids": r.evidence_ids,
            })

    def _ownership_key(row: dict) -> tuple[str, str, str, str]:
        """把 SRM 同一页面的重复展示（万/万元、0.99/99.00%）归为一条。"""
        def amount(value: object) -> str:
            return re.sub(r"\s|万元?|人民币|元", "", str(value or "")).lower()
        def ratio(value: object) -> str:
            original = str(value or "")
            text = re.sub(r"\s|%", "", original)
            try:
                number = float(text)
                if "%" not in original and number <= 1:
                    number *= 100
                return f"{number:.6f}"
            except ValueError:
                return text
        return (
            str(row.get("supplier") or ""), str(row.get("shareholder") or ""),
            amount(row.get("amount")), ratio(row.get("ratio")),
        )

    unique_ownership: dict[tuple[str, str, str, str, str], dict] = {}
    for row in ownership_rows_raw:
        key = _ownership_key(row)
        previous = unique_ownership.get(key)
        if previous is None:
            unique_ownership[key] = row
        else:
            previous["evidence_ids"] = sorted(set(previous.get("evidence_ids", [])) | set(row.get("evidence_ids", [])))
    ownership_rows_raw = list(unique_ownership.values())

    def _section_rows(record_kind: str, labels: dict[str, str]) -> list[dict]:
        out = []
        for r in ext.records:
            if r.record_kind != record_kind:
                continue
            out.append({
                "supplier": names.get(r.supplier_id or "", r.subject_name or "未归属"),
                "fields": {label: r.fields.get(key, "未返回") for key, label in labels.items()},
                "evidence_ids": r.evidence_ids,
            })
        return out

    srm_branches = _section_rows("branch", {
        "企业名称": "企业名称", "成立日期": "成立日期", "企业状态": "企业状态", "法人": "法人",
    })
    srm_personnel = _section_rows("personnel", {"姓名": "姓名", "职位": "职位"})

    def _duplicate_result(rows: list[dict], fields: list[str], empty: str = "未取得，无法进行重复校验") -> str:
        values: dict[str, set[str]] = {}
        for row in rows:
            supplier = str(row.get("supplier") or "未归属")
            source = row.get("fields") or row
            for field in fields:
                value = str(source.get(field) or "").strip()
                if value and value not in {"未返回", "未取得", "—"}:
                    values.setdefault(f"{field}:{value}", set()).add(supplier)
        duplicates = [key for key, owners in values.items() if len(owners) > 1]
        if duplicates:
            return "发现重复：" + "；".join(duplicates[:8])
        if not values:
            return empty
        return "未发现重复"

    registration_for_check = [r for r in srm_registration if r["fields"]]
    srm_duplicate_checks = {
        "工商信息": _duplicate_result(registration_for_check, ["企业名称", "统一社会信用代码", "法定代表人"]),
        "股东信息": _duplicate_result(ownership_rows_raw, ["股东名称", "股东", "shareholder", "ratio", "认缴比例", "持股比例"]),
        "分支机构": _duplicate_result(srm_branches, ["企业名称", "法人"]),
        "主要人员": _duplicate_result(srm_personnel, ["姓名", "职位"]),
    }

    # SRM 工商信息统一为“横轴公司、纵轴统计项目”，避免 16 列横向表格在 PDF 中挤压、错位。
    registration_labels = [
        ("企业名称", ("企业名称",)),
        ("法定代表人", ("法定代表人", "法人代表")),
        ("注册资本", ("注册资本",)),
        ("实缴资本", ("实缴资本",)),
        ("成立日期", ("成立日期",)),
        ("企业状态", ("企业状态",)),
        ("统一社会信用代码", ("统一社会信用代码", "信用代码")),
        ("企业类型", ("企业类型",)),
        ("行业", ("一级行业", "行业大类", "二级行业", "三级行业", "四级行业")),
        ("营业期限", ("营业期限",)),
        ("参保人数", ("参保人数",)),
        ("曾用名", ("曾用名",)),
        ("注册地址", ("注册地址", "地址")),
    ]
    registration_by_supplier = {
        r["supplier"]: r.get("fields") or {}
        for r in srm_registration
        if r.get("fields")
    }
    registration_matrix = {
        "columns": [names.get(s.supplier_id, _display_company_name(s.display_name)) for s in sups],
        "rows": [],
    }
    for label, aliases in registration_labels:
        cells = []
        has_value = False
        for s in sups:
            fields = registration_by_supplier.get(names.get(s.supplier_id, _display_company_name(s.display_name)), {})
            value = _first_field_value(fields, aliases) or "未取得"
            cells.append(value)
            has_value = has_value or value != "未取得"
        # 只保留至少有一家公司真正有数据的统计项，避免报告出现空行。
        if has_value:
            registration_matrix["rows"].append({"label": label, "cells": cells})
    government_screenshots = []
    for ev in evidence.values():
        if ev.field == "external.government_procurement":
            loc = ev.location
            if loc.get("kind") in ("snapshot", "url"):
                ref = loc.get("file") or loc.get("path") or loc.get("url") or "未提供"
                image = None
                local_file = loc.get("file")
                if local_file and Path(project_dir / str(local_file)).suffix.lower() in {".png", ".jpg", ".jpeg"}:
                    source = project_dir / str(local_file)
                    if source.exists():
                        image = os.path.relpath(source, project_dir / "output")
                government_screenshots.append({"label": "中国政府采购网查询证据", "ref": ref, "image": image})
    coverage_map = {(c.supplier_id, c.source_id): c.status for c in findings_file.coverage}
    sources = ["government_procurement", "srm"]

    # 摘要
    supplier_rows = []
    unfinished_by_sup: dict[str, int] = {}
    for c in findings_file.coverage:
        if c.status not in ("match", "no_match_verified"):
            unfinished_by_sup[c.supplier_id] = unfinished_by_sup.get(c.supplier_id, 0) + 1
    review_set = set(findings_file.human_review_queue)
    for s in sups:
        rows = [f for f in findings if s.supplier_id in f.supplier_ids]
        supplier_rows.append({
            "name": names.get(s.supplier_id, _display_company_name(s.display_name)),
            "i": sum(1 for f in rows if f.level == "I"),
            "ii": sum(1 for f in rows if f.level == "II"),
            "iii": sum(1 for f in rows if f.level == "III"),
            "review": sum(1 for f in rows if f.finding_id in review_set),
            "unfinished": unfinished_by_sup.get(s.supplier_id, 0),
            "srm_status": _query_status_label(coverage_map.get((s.supplier_id, "srm"), "not_queried")),
            "gov_status": _query_status_label(coverage_map.get((s.supplier_id, "government_procurement"), "not_queried")),
        })
    coverage_matrix = [
        {
            "supplier": names.get(s.supplier_id, _display_company_name(s.display_name)),
            "cells": [_query_status_label(coverage_map.get((s.supplier_id, sid), "not_queried")) for sid in sources],
        }
        for s in sups
    ]
    coverage_sep = "|---" * (1 + len(sources)) + "|"
    headline = [_frow(f, names) for f in findings if f.level == "I"]

    # 公司信息交叉表
    def contacts_of(sup_id: str, kind: str) -> list[str]:
        return sorted({c.raw_masked for c in entities.contacts if c.supplier_id == sup_id and c.kind == kind}) or ["—"]

    def parties_of(sup_id: str, role: str) -> list[str]:
        values = sorted({p.name for p in entities.parties if p.supplier_id == sup_id and p.role == role})
        if len(values) <= 1:
            return values or ["—"]
        return ["待人工复核（多个候选）"]

    cross_tables = []
    if sups:
        cols = [names.get(s.supplier_id, _display_company_name(s.display_name)) for s in sups]

        def table(title: str, getter) -> dict:
            """行=值（共同值标注），列=供应商，命中打勾。"""
            seen: dict[str, list] = {}
            for s in sups:
                for v in getter(s.supplier_id):
                    if v == "—":
                        continue  # 占位符（无数据）不得显示为共同值
                    seen.setdefault(v, []).append(names.get(s.supplier_id, _display_company_name(s.display_name)))
            rows = []
            for idx, (v, owners) in enumerate(sorted(seen.items(), key=lambda kv: (-len(kv[1]), kv[0])), start=1):
                rows.append({
                    "label": f"{idx}. {v}" + ("（共同）" if len(owners) > 1 else ""),
                    "cells": ["✔" if names.get(s.supplier_id, _display_company_name(s.display_name)) in owners else "" for s in sups],
                })
            return {"title": title, "columns": cols, "rows": rows,
                    "sep": "|---" * (1 + len(cols)) + "|"}

        cross_tables.append({
            "title": "统一社会信用代码",
            "columns": cols,
            "rows": [{
                "label": f"{idx}. {names.get(s.supplier_id, _display_company_name(s.display_name))}：{s.uscc or '（未提取）'}"
                         + ("" if s.uscc_status == "present_valid" else f"［{s.uscc_status}］"),
                "cells": ["✔" if c is s else "" for c in sups],
            } for idx, s in enumerate(sups, start=1)],
            "sep": "|---" * (1 + len(cols)) + "|",
        })
        cross_tables.append(table("联系电话", lambda sid: contacts_of(sid, "phone")))
        cross_tables.append(table("电子邮箱", lambda sid: contacts_of(sid, "email")))
        cross_tables.append(table("地址", lambda sid: contacts_of(sid, "address")))
        cross_tables.append(table("法定代表人", lambda sid: parties_of(sid, "legal_rep")))
        cross_tables.append(table("授权代表", lambda sid: parties_of(sid, "bid_agent")))
        cross_tables.append(table("股东", lambda sid: parties_of(sid, "shareholder")))
        cross_tables.append(table("联系人", lambda sid: parties_of(sid, "contact")))

    # 分域发现
    metadata_findings = [_frow(f, names) for f in findings if f.domain == "metadata"]
    ownership_judicial = [_frow(f, names) for f in findings if f.domain in ("ownership", "judicial")]
    dishonesty_findings = [_frow(f, names) for f in findings if f.domain == "dishonesty"]
    review_findings = [_frow(f, names) for f in findings if f.finding_id in review_set or f.status == "human_review_required"]

    # 股权
    ownership_rows = [
        {
            "relation_id": o.relation_id,
            "from_label": o.from_party_name or o.from_company_name or o.from_company_uscc or "—",
            "to_company_name": o.to_company_name + (f"（供应商 {names.get(o.supplier_id, o.supplier_id)}）" if o.supplier_id and o.supplier_id in names else ""),
            "share_ratio": o.share_ratio,
            "relation_date": str(o.relation_date) if o.relation_date else None,
            "relation_end_date": str(o.relation_end_date) if o.relation_end_date else None,
            "confirmation_label": {"confirmed": "已确认", "candidate": "候选", "unconfirmed": "未确认"}[o.subject_confirmation],
            "evidence_ids": o.evidence_ids,
        }
        for o in ext.ownership
    ]

    # 司法记录
    judicial_rows = []
    for r in ext.records:
        if r.record_kind != "judicial_case":
            continue
        judicial_rows.append({
            "supplier": names.get(r.supplier_id or "", "（未归属）"),
            "record_id": r.record_id,
            "confirmation_label": {"confirmed": "已确认", "candidate": "候选", "unconfirmed": "未确认"}[r.subject_confirmation],
            "desc": "；".join(f"{k}={v}" for k, v in list(r.fields.items())[:5]) or r.subject_name or "",
            "evidence_ids": r.evidence_ids,
        })

    # 违法失信按渠道
    dishonesty_sources = []
    for sid, label in SOURCE_LABELS.items():
        status_rows = [
            {"line": f"{names.get(s.supplier_id, _display_company_name(s.display_name))}：{_query_status_label(coverage_map.get((s.supplier_id, sid), 'not_queried'))}"}
            for s in sups
        ]
        records = []
        for r in ext.records:
            if r.source_id != sid or r.record_kind not in ("dishonesty", "penalty"):
                continue
            f = r.fields
            records.append({
                "supplier": names.get(r.supplier_id or "", "（未归属）"),
                "behavior": str(f.get("行为") or f.get("title") or f.get("visible_result") or f.get("当事人") or "—"),
                "penalty": str(f.get("处罚") or f.get("处理结果") or "—"),
                "reference": str(f.get("依据") or f.get("处罚决定文号") or f.get("处理编号") or f.get("决定书文号") or "—"),
                "date": str(f.get("处理日期") or f.get("处罚日期") or f.get("发布日期") or "—"),
                "validity": (f"{r.effective_from or '未知'} ~ {r.effective_to or '未载明'}") if (r.effective_from or r.effective_to) else "未载明",
                "status": "match" if (r.subject_confirmation == "confirmed") else r.subject_confirmation,
                "evidence_ids": r.evidence_ids,
            })
        other_kinds = [r for r in ext.records if r.source_id == sid and r.record_kind not in ("dishonesty", "penalty")]
        dishonesty_sources.append({
            "label": label, "status_rows": status_rows, "records": records,
            "other_note": ("该渠道的司法/股权类记录见第 5 章。" if (not records and other_kinds) else None),
        })

    # 复核缺口
    gaps: list[str] = []
    for a in inventory.anomalies:
        gaps.append(f"清单异常：{a.detail or a.anomaly}")
    for a in content.anomalies:
        gaps.append(f"提取异常：{a.detail or a.anomaly}")
    if low_conf_ids:
        gaps.append(f"低置信度字段（如 OCR）{len(low_conf_ids)} 项未参与匹配，需人工核对（见人工复核清单.csv）")
    for c in findings_file.coverage:
        if c.status not in ("match", "no_match_verified"):
            gaps.append(f"查询覆盖：{names.get(c.supplier_id, c.supplier_id)} × {SOURCE_LABELS.get(c.source_id, c.source_id)}：{_query_status_label(c.status)}")
    for s in sups:
        if s.uscc_status != "present_valid":
            gaps.append(f"主体确认：{names.get(s.supplier_id, _display_company_name(s.display_name))} 统一社会信用代码状态为 {s.uscc_status}，相关外部记录暂不归属")
    if cfg.procurement_rules_source in (None, "", "unspecified"):
        gaps.append("采购文件资格/否决条款未提供：失信记录的资格影响只能标注“待采购人确认”")
    if cfg.bid_deadline is None:
        gaps.append("project.yaml 未提供 bid_deadline：禁止产生“处罚处于有效期内”的结论")

    # 附录
    queries_rows = [
        {
            "source": SOURCE_LABELS.get(q.source_id, q.source_id),
            "supplier": names.get(q.subject_supplier_id or "", "（未归属）"),
            "status": _query_status_label(q.status),
            "mode": {"browser_session": "浏览器会话", "official_api": "官方接口", "public_web": "公开网页", "manual_import": "人工导入", "none": "未执行"}.get(q.query_mode, "未取得"),
            "queried_at": str(q.queried_at) if q.queried_at else None,
            "detail": (q.detail or "")[:120],
        }
        for q in sorted(ext.queries, key=lambda x: (x.source_id, x.subject_supplier_id or ""))
    ]
    appendix = {
        "documents": [
            {
                "relative_path": d.relative_path,
                "category": d.category,
                "size_bytes": d.size_bytes,
                "short_sha": d.sha256[:16],
                "status": d.extraction_status,
                "detail": d.status_detail,
            }
            for d in inventory.documents
        ],
        "rule_ids": sorted({f.rule_id for f in findings}),
        "queries": queries_rows,
    }

    deadline_display = (
        "未提供（无法判断处罚/禁入是否处于有效期）"
        if cfg.bid_deadline is None
        else f"{cfg.bid_deadline.date().isoformat()}（投标截止时间）"
    )

    return {
        "project": cfg,
        "project_dir": project_dir,
        "run": inventory.run,
        "project_identity": project_identity,
        "business_rows": business_rows,
        "srm_registration": srm_registration,
        "srm_registration_matrix": registration_matrix,
        "srm_ownership": ownership_rows_raw,
        "srm_branches": srm_branches,
        "srm_personnel": srm_personnel,
        "srm_duplicate_checks": srm_duplicate_checks,
        "srm_query_status": {names.get(s.supplier_id, _display_company_name(s.display_name)): _query_status_label(_srm_query_status(s.supplier_id)) for s in sups},
        "government_screenshots": government_screenshots,
        "summary": {
            "deadline_display": deadline_display,
            "supplier_count": len(sups),
            "supplier_names": [names.get(s.supplier_id, _display_company_name(s.display_name)) for s in sups],
            "bid_docs": sum(1 for d in inventory.documents if d.category == "bid"),
            "proc_docs": sum(1 for d in inventory.documents if d.category == "procurement"),
            "ext_dir": sum(1 for d in inventory.documents if d.category == "external_evidence"),
            "supplier_rows": supplier_rows,
            "coverage_matrix": coverage_matrix,
            "coverage_sep": coverage_sep,
            "coverage_sources": [SOURCE_LABELS.get(s, s) for s in sources],
            "coverage_gaps": sum(unfinished_by_sup.values()),
            "headline_findings": headline,
        },
        "coverage_sources": [SOURCE_LABELS.get(s, s) for s in sources],
        "suppliers": sups,
        "cross_tables": cross_tables,
        "metadata_findings": metadata_findings,
        "ownership": ownership_rows,
        "judicial_rows": judicial_rows,
        "ownership_judicial_findings": ownership_judicial,
        "dishonesty_sources": dishonesty_sources,
        "dishonesty_findings": dishonesty_findings,
        "review_findings": review_findings,
        "review_gaps": gaps,
        "quote_rows": [
            {
                "supplier": names.get(s.supplier_id, _display_company_name(s.display_name)),
                "total": (f"{quote_by_supplier[s.supplier_id]['total']:.2f}" if quote_by_supplier.get(s.supplier_id, {}).get("total") is not None else "未取得"),
                "components": "、".join(f"{x['name']} {x['amount']:.2f}" for x in quote_by_supplier.get(s.supplier_id, {}).get("components", [])) or "未取得",
                "arithmetic_status": quote_by_supplier.get(s.supplier_id, {}).get("arithmetic_status", "未取得报价分析结构化结果"),
                "anomaly_status": quote_by_supplier.get(s.supplier_id, {}).get("anomaly_status", "未取得报价分析结构化结果"),
                "evidence_ids": "、".join(quote_by_supplier.get(s.supplier_id, {}).get("evidence_ids", [])) or "—",
            }
            for s in sups
        ],
        "appendix": appendix,
    }


# --------------------------------------------------------------------- 渲染


def _render_markdown(ctx: dict, path: Path) -> None:
    from jinja2 import Environment, FileSystemLoader

    # 报告表格依赖模板循环保留换行；不能用 trim_blocks 把表头、分隔线和数据
    # 拼成一行，否则 Markdown/PDF 会把纵向对照表解析成普通段落。
    env = Environment(loader=FileSystemLoader(str(TEMPLATE_PATH.parent)), autoescape=False,
                      trim_blocks=False, lstrip_blocks=False)
    tpl = env.get_template(TEMPLATE_PATH.name)
    text = tpl.render(**ctx)
    path.write_text(text, encoding="utf-8")


def _render_evidence_csv(evidence: dict[str, Evidence], inventory: InventoryFile, path: Path) -> None:
    doc_path = {d.document_id: d.relative_path for d in inventory.documents}
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["evidence_id", "source_type", "document_path", "location", "field",
                    "raw_value_masked", "normalized", "method", "confidence", "strength",
                    "collected_at", "source_sha256", "note"])
        for ev in sorted(evidence.values(), key=lambda x: x.evidence_id):
            w.writerow([
                ev.evidence_id, ev.source_type,
                doc_path.get(ev.document_id or "", ev.document_id or ""),
                _loc_str(ev),
                ev.field, ev.raw_value, ev.normalized_value or "", ev.method,
                ev.confidence if ev.confidence is not None else "", ev.strength,
                ev.collected_at.isoformat(), ev.sha256 or "", ev.note or "",
            ])


def _loc_str(ev: Evidence) -> str:
    parts = []
    for k, v in ev.location.items():
        if v is None:
            continue
        parts.append(f"{k}={v}")
    return " ".join(parts)[:300]


def _render_review_csv(findings_file, entities, low_conf_ids, inventory, content, path: Path,
                       field_decisions: list[dict] | None = None) -> None:
    names = _sup_name(entities)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["类别", "编号", "预警级别", "涉及供应商", "内容", "规则/原因", "证据强度", "建议动作"])
        for f in findings_file.findings:
            if f.status == "human_review_required":
                w.writerow(["风险发现", f.finding_id, f.level,
                            "、".join(names.get(s, s) for s in f.supplier_ids),
                            f.fact, f.rule_id, f.evidence_strength, f.recommendation])
        for eid in sorted(low_conf_ids):
            w.writerow(["低置信度字段", eid, "—", "—", "OCR/低置信度提取结果未参与匹配，需人工核对原件", "低置信度", "D", "对照原件核对字段值"])
        field_labels = {
            "project_name": "项目名称", "tenderer": "招标人", "project_code": "项目编号",
            "bid_date": "投标日期", "company_name": "投标人", "uscc": "统一社会信用代码",
            "legal_rep_name": "法定代表人", "legal_rep_id": "法定代表人身份证号",
            "bid_agent_name": "授权代表", "bid_agent_id": "授权代表身份证号",
        }
        for decision in field_decisions or []:
            if decision.get("status") == "selected":
                continue
            candidates = decision.get("candidate_values") or []
            candidate_text = "、".join(str(x.get("value") or "") for x in candidates[:8]) or "未取得"
            scope = str(decision.get("scope") or "")
            supplier = scope.removeprefix("supplier:") if scope.startswith("supplier:") else "项目"
            w.writerow([
                "字段冲突/缺口", f"FIELD-{decision.get('scope')}-{decision.get('field')}", "—", supplier,
                f"{field_labels.get(decision.get('field'), decision.get('field'))}：{candidate_text}；未进入正式报告单值字段",
                decision.get("reason") or decision.get("status"), "—", "核对原件、页码/坐标和标签关系后确认唯一值",
            ])
        for a in inventory.anomalies:
            w.writerow(["清单异常", "—", "—", "—", a.detail or a.anomaly, a.anomaly, "—", "人工检查文件并补录"])
        for d in inventory.documents:
            if d.extraction_status not in ("ok", "partial", "skipped"):
                w.writerow(["读取失败文件", "—", "—", "—",
                            f"{d.relative_path}：{d.extraction_status}（{d.status_detail or ''}）",
                            d.extraction_status, "—", "人工检查文件并补录完整版本"])
        for a in content.anomalies:
            w.writerow(["提取异常", "—", "—", a.relative_path or "—",
                        f"{a.relative_path or ''}：{a.detail or a.anomaly}".strip("："),
                        a.anomaly, "—", "人工核对文件内容"])


def _render_docx(ctx: dict, path: Path) -> None:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    d.add_heading(f"投标清标报告：{ctx['project'].project_name}", level=0)
    p = d.add_paragraph(
        f"项目编号 {ctx['project'].project_id}；投标截止时间 {ctx['summary']['deadline_display']}；"
        f"运行标识 {ctx['run'].run_id}。本报告为风险线索整理结果，非最终认定。"
    )
    p.runs[0].font.size = Pt(9)
    d.add_heading("1. 封面信息页", level=1)
    d.add_paragraph("招标人：" + "；".join(ctx["project_identity"]["tenderer"] or ["未取得"]))
    d.add_paragraph("项目名称：" + "；".join(ctx["project_identity"]["project_name"] or ["未取得"]))
    d.add_paragraph("投标人：" + "、".join(ctx["summary"]["supplier_names"]))
    d.add_heading("2. 基本信息页", level=1)
    for r in ctx["business_rows"]:
        d.add_paragraph(
            f"{r['supplier']}：{r['company']}；统一社会信用代码 {r['uscc']}；"
            f"法定代表人 {r['legal_name']}（{r['legal_id']}）；"
            f"授权代表 {r['agent_name']}（{r['agent_id']}）"
        )
    d.add_heading("3. 工商信息", level=1)
    for r in ctx["srm_registration"]:
        d.add_paragraph(f"{r['supplier']}：{r['fields'] or '登录成功但未取得工商记录'}")
    d.add_paragraph("交叉校验：" + ctx["srm_duplicate_checks"]["工商信息"])
    d.add_heading("4. 股东信息", level=1)
    for r in ctx["srm_ownership"]:
        d.add_paragraph(f"{r['supplier']}：{r['shareholder']}；投资金额 {r['amount']}；认缴比例 {r['ratio']}")
    d.add_paragraph("交叉校验：" + ctx["srm_duplicate_checks"]["股东信息"])
    d.add_heading("5. 分支机构", level=1)
    for r in ctx["srm_branches"]:
        d.add_paragraph(f"{r['supplier']}：{r['fields']}")
    d.add_paragraph("交叉校验：" + ctx["srm_duplicate_checks"]["分支机构"])
    d.add_heading("6. 主要人员", level=1)
    for r in ctx["srm_personnel"]:
        d.add_paragraph(f"{r['supplier']}：{r['fields']}")
    d.add_paragraph("交叉校验：" + ctx["srm_duplicate_checks"]["主要人员"])
    d.add_heading("7. 政府采购网失信信息及证据截图", level=1)
    for src in ctx["dishonesty_sources"]:
        if src["label"] != "中国政府采购网":
            continue
        d.add_paragraph("；".join(r["line"] for r in src["status_rows"]))
        for r in src["records"]:
            d.add_paragraph(f"{r['supplier']}：{r['behavior']} / {r['penalty']}（{r['date']}）", style="List Bullet")
    for shot in ctx["government_screenshots"]:
        image_ref = shot.get("image")
        if image_ref:
            try:
                from docx.shared import Inches
                d.add_paragraph(f"证据截图：{shot['ref']}")
                d.add_picture(str(ctx["project_dir"] / "output" / image_ref), width=Inches(6.2))
            except Exception:  # noqa: BLE001 - 图片损坏时不阻断文字报告
                d.add_paragraph(f"证据截图：{shot['ref']}（图片无法嵌入，请按引用核对）")
    d.add_heading("8. PDF、Office 与扫描件文件线索", level=1)
    for f in ctx["metadata_findings"]:
        d.add_paragraph(f"[{f['level_label']}|{f['evidence_strength']}] {f['fact']}（{f['finding_id']}）")
    d.add_heading("9. 清标结论与三级预警", level=1)
    for s in ctx["summary"]["supplier_rows"]:
        d.add_paragraph(
            f"{s['name']}：I 级 {s['i']} 项；II 级 {s['ii']} 项；III 级 {s['iii']} 项；"
            f"必须人工复核 {s['review']} 项。"
        )
    for f in ctx["summary"]["headline_findings"]:
        d.add_paragraph(f"{f['finding_id']}（{f['rule_id']}）：{f['fact']}", style="List Bullet")
    d.add_heading("10. 证据与人工复核", level=1)
    for f in ctx["review_findings"]:
        d.add_paragraph(f"{f['finding_id']}：{f['fact']}", style="List Bullet")
    for g in ctx["review_gaps"]:
        d.add_paragraph(g, style="List Bullet")
    d.add_heading("附录：输入文件清单", level=1)
    t = d.add_table(rows=1, cols=4)
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "文件", "类别", "状态", "SHA-256（前16位）"
    for doc in ctx["appendix"]["documents"]:
        cells = t.add_row().cells
        cells[0].text = doc["relative_path"]
        cells[1].text = doc["category"]
        cells[2].text = doc["status"]
        cells[3].text = doc["short_sha"]
    d.save(str(path))


if __name__ == "__main__":
    app()
