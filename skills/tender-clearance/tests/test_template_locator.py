"""Word 模板解析与解析结果安全边界。"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from tc.template_locator import load_template_spec, parse_template, resolve_template


def _make_template(path: Path) -> None:
    document = Document()
    document.add_heading("技术要求", level=1)
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "指标名称"
    table.rows[0].cells[1].text = "技术要求"
    table.rows[0].cells[2].text = "响应"
    row = table.add_row().cells
    row[0].text = "额定功率（kW）"
    row[1].text = ""
    row[2].text = ""
    row = table.add_row().cells
    row[0].text = "防护等级"
    row[1].text = ""
    row[2].text = ""
    document.save(path)


def test_parse_template_extracts_stable_metric_anchors(tmp_path):
    path = tmp_path / "空白招标文件模板.docx"
    _make_template(path)

    first = parse_template(path)
    second = parse_template(path)

    assert first.mode == "template_aligned"
    assert [item.label for item in first.metrics] == ["额定功率", "防护等级"]
    assert first.metrics[0].unit == "kW"
    assert [item.metric_id for item in first.metrics] == [item.metric_id for item in second.metrics]
    assert first.sections == ["技术要求"]


def test_template_is_optional_and_ambiguous_candidates_are_not_guessed(tmp_path):
    (tmp_path / "procurement").mkdir()
    empty = resolve_template(tmp_path)
    assert empty.status == "not_provided"
    assert "大幅提高" in empty.message

    _make_template(tmp_path / "procurement" / "招标模板A.docx")
    _make_template(tmp_path / "procurement" / "招标模板B.docx")
    ambiguous = resolve_template(tmp_path)
    assert ambiguous.status == "ambiguous"
    assert len(ambiguous.candidates) == 2


def test_load_template_spec_keeps_generic_fallback_when_not_provided(tmp_path):
    spec = load_template_spec(tmp_path)
    assert spec.mode == "generic_ocr"
    assert spec.metrics == []
    assert "通用 OCR 兜底" in spec.message
